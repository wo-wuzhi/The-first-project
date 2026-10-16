from docx import Document

path = r"D:\edge\接口文档.docx"
doc = Document(path)

print("===== PARAGRAPHS =====")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print(f"[{i}] {t}")

print()
print("===== TABLES =====")
for ti, table in enumerate(doc.tables):
    print(f"--- TABLE {ti} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        print(f"  R{ri}: {cells}")
