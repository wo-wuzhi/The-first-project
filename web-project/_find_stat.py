from docx import Document
doc = Document(r"D:\edge\接口文档.docx")

print("=== A. 含 统计/人数/报表 的段落 ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t and any(k in t for k in ("统计", "人数", "报表")):
        print(f"[{i}] {t}")

print()
print("=== B. 表格总数 ===", len(doc.tables))

print()
print("=== C. 每个表格首行（定位用）===")
for ti, tb in enumerate(doc.tables):
    first = " | ".join(c.text.strip().replace("\n", "/")[:50] for c in tb.rows[0].cells)
    print(f"T{ti}: {first}")
