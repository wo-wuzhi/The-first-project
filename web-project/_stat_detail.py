from docx import Document
doc = Document(r"D:\edge\接口文档.docx")

print("=== 段落 460-560 ===")
for i, p in enumerate(doc.paragraphs):
    if 460 <= i <= 560:
        t = p.text.strip()
        if t:
            print(f"[{i}] {t}")

print()
print("=== 表格 T30-T40 ===")
for ti in range(30, min(41, len(doc.tables))):
    tb = doc.tables[ti]
    print(f"\n--- T{ti} ({len(tb.rows)}x{len(tb.columns)}) ---")
    for ri, row in enumerate(tb.rows):
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        print(f"  R{ri}: {cells}")
